"""
Einmal/Merge: Excel RKRS+ Einsatztauglichkeit -> pilotenliste.json
Schriftfarbe Spalte A (Vorname): rot=1, schwarz=2, grün=4 (erste_erfahrung).

Zielpfad wie rkrk_assess_live_personel.php:
  $rkrkLiveassessRoot = $folderoffset . $RKRKRelRKRKDataLocation;
  $rkrkLiveassessData = $rkrkLiveassessRoot . "data/liveassess";
  $rkrkLivePersonelRoot = $rkrkLiveassessData . "/personel";
  → pilotenliste.json liegt in $rkrkLivePersonelRoot.

$RKRKRelRKRKDataLocation kommt aus scripts/rkrk_gen.php — bei Pfadänderung dort
dieses Skript anpassen (Konstanten unten).
"""
from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

# Wie rkrk_gen.php (include aus assess/: $folderoffset = "../")
_FOLDEROFFSET = "../"
_RKRK_REL_MISSION_LOCATION = "../../rkropen/"


def _rkrk_fix_path(p: str) -> str:
    p = str(p).replace("\\", "/")
    while "/rkropen/rkropen/" in p:
        p = p.replace("/rkropen/rkropen/", "/rkropen/")
    return p


def _php_rkrk_rel_data_location() -> str:
    """Entspricht $RKRKRelRKRKDataLocation in rkrk_gen.php."""
    return _rkrk_fix_path(_RKRK_REL_MISSION_LOCATION + "rkrk/")


def _default_live_personel_dir() -> Path:
    """
    Entspricht $rkrkLivePersonelRoot relativ zum assess/-Ordner dieses Skripts:
    ($folderoffset . $RKRKRelRKRKDataLocation) + data/liveassess/personel
    """
    assess_dir = Path(__file__).resolve().parent
    rkrk_liveassess_root = (assess_dir / (_FOLDEROFFSET + _php_rkrk_rel_data_location())).resolve()
    rkrk_liveassess_data = rkrk_liveassess_root / "data" / "liveassess"
    return rkrk_liveassess_data / "personel"

try:
    from openpyxl import load_workbook
except ImportError:
    print("pip install openpyxl", file=sys.stderr)
    raise

# Excel: lokaler Pfad oder TEMP-Kopie (OneDrive blockiert evtl. direktes Lesen)
DEFAULT_XLSX = Path(r"c:\Users\tchr7\OneDrive\Rehkitzrettung\RKRS PLUS\20250919 Einsatztauglichkeit.xlsx")
TEMP_COPY = Path(os.environ.get("TEMP", ".")) / "einsatztauglichkeit_merge.xlsx"


def pick_xlsx() -> Path:
    if len(sys.argv) > 1:
        return Path(sys.argv[1])
    if DEFAULT_XLSX.is_file():
        return DEFAULT_XLSX
    if TEMP_COPY.is_file():
        return TEMP_COPY
    sys.exit(f"Keine Excel-Datei: {DEFAULT_XLSX} oder {TEMP_COPY}")


def font_rgb_upper(cell) -> str | None:
    c = cell.font.color
    if c is None:
        return None
    rgb = getattr(c, "rgb", None)
    if rgb:
        return str(rgb).upper()
    return None


def erste_erfahrung_from_vorname_cell(rgb: str | None) -> int | None:
    """rot=1, schwarz=2, grün=4."""
    if not rgb:
        return None
    # ARGB 8 hex
    if re.search(r"FF0000$", rgb) or rgb.endswith("FF0000"):  # red
        return 1
    if re.search(r"000000$", rgb) or rgb.endswith("000000"):  # black
        return 2
    if "00B050" in rgb or rgb.endswith("00B050"):  # Excel green
        return 4
    return None


def norm_part(s: str) -> str:
    return " ".join(s.lower().split())


def row_key(vor: str, nach: str, email: str) -> str:
    e = email.strip().lower()
    if e:
        return "e:" + e
    return "n:" + norm_part(vor) + "|" + norm_part(nach)


def main() -> None:
    xlsx = pick_xlsx()
    wb_styles = load_workbook(xlsx, data_only=False)
    wb_vals = load_workbook(xlsx, data_only=True)
    ws_s = wb_styles.active
    ws_v = wb_vals.active

    merged: dict[str, dict] = {}
    excel_order: list[str] = []
    seen_excel: set[str] = set()

    for ri in range(2, ws_v.max_row + 1):
        vor_c = ws_v.cell(ri, 1).value
        nach_c = ws_v.cell(ri, 2).value
        vor = "" if vor_c is None else str(vor_c).strip()
        nach = "" if nach_c is None else str(nach_c).strip()
        if not vor and not nach:
            continue

        strasse_c = ws_v.cell(ri, 3).value
        plz_c = ws_v.cell(ri, 4).value
        ort_c = ws_v.cell(ri, 5).value
        email_c = ws_v.cell(ri, 6).value
        tel_c = ws_v.cell(ri, 7).value

        strasse = "" if strasse_c is None else str(strasse_c).strip()
        plz = "" if plz_c is None else str(plz_c).strip()
        ort = "" if ort_c is None else str(ort_c).strip()
        email = "" if email_c is None else str(email_c).strip()
        telefon = "" if tel_c is None else str(tel_c).strip()

        rgb = font_rgb_upper(ws_s.cell(ri, 1))
        ee = erste_erfahrung_from_vorname_cell(rgb)
        if ee is None:
            print(f"Warnung Zeile {ri}: unbekannte Schriftfarbe {rgb!r} ({vor} {nach})", file=sys.stderr)

        display = (vor + " " + nach).strip()
        row = {
            "vorname": vor,
            "nachname": nach,
            "name": display,
            "strasse": strasse,
            "plz": plz,
            "ort": ort,
            "email": email,
            "telefon": telefon,
        }
        if ee is not None:
            row["erste_erfahrung"] = ee
        k = row_key(vor, nach, email)
        merged[k] = row
        if k not in seen_excel:
            excel_order.append(k)
            seen_excel.add(k)

    # Ausgabe: gleicher Ordner wie PHP $rkrkLivePersonelRoot (oder RKRK_PILOT_OUT=Verzeichnis überschreiben)
    _default_personel = _default_live_personel_dir()
    out_dir = Path(os.environ.get("RKRK_PILOT_OUT", str(_default_personel)))
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "pilotenliste.json"

    existing_pilots: list[dict] = []
    if out_path.is_file():
        raw = json.loads(out_path.read_text(encoding="utf-8"))
        if isinstance(raw, dict) and isinstance(raw.get("pilots"), list):
            existing_pilots = [p for p in raw["pilots"] if isinstance(p, dict)]

    by_key: dict[str, dict] = {k: dict(v) for k, v in merged.items()}

    pilots_out: list[dict] = []
    used: set[str] = set()

    for p in existing_pilots:
        vor = str(p.get("vorname", "")).strip()
        nach = str(p.get("nachname", "")).strip()
        em = str(p.get("email", "")).strip()
        k = row_key(vor, nach, em)
        base = dict(p)
        if k in merged:
            base.update(merged[k])
        pilots_out.append(base)
        used.add(k)

    for k in excel_order:
        if k in used:
            continue
        pilots_out.append(by_key[k])
        used.add(k)

    # Spaltenreihenfolge
    base_cols = [
        "vorname",
        "nachname",
        "name",
        "strasse",
        "plz",
        "ort",
        "email",
        "telefon",
        "erste_erfahrung",
    ]
    extra: set[str] = set()
    for p in pilots_out:
        for key in p:
            if key != "id":
                extra.add(key)
    columns = [c for c in base_cols if c in extra]
    for c in sorted(extra - set(columns)):
        if c != "id":
            columns.append(c)

    payload = {"columns": columns, "pilots": pilots_out}
    out_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"OK {len(pilots_out)} Piloten -> {out_path}")


if __name__ == "__main__":
    main()

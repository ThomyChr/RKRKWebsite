import argparse
import json
import os
import secrets
import shutil
import sys
import time
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from uuid import uuid4

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Missing dependency: openpyxl. Install via pip install openpyxl") from exc


EXCEL_DATE_BASE = date(1899, 12, 30)

SHEET_TRANSACTION = "KB CS Rehkitzrettung"
SHEET_BUDGET = "Budget"
SHEET_INVENTORY = "Inventur"
SHEET_COST_CATEGORIES = "Kostenarten"

TRANSACTION_HEADER_ROW = 13
BUDGET_HEADER_ROW = 4
INVENTORY_HEADER_ROW = 3
COST_CATEGORY_HEADER_ROW = 1


@dataclass
class Issue:
    sheet_name: str
    row_no: Optional[int]
    column_name: Optional[str]
    severity: str
    issue_code: str
    message: str
    raw_value: Optional[str]


class Importer:
    def __init__(self, workbook_path: str, output_dir: str, dry_run: bool = True):
        self.workbook_path = workbook_path
        self.output_dir = Path(output_dir)
        self.dry_run = dry_run
        self.issues: List[Issue] = []
        self.stats = {
            "cost_categories": 0,
            "transactions": 0,
            "budget_lines": 0,
            "inventory_items": 0,
            "warnings": 0,
            "errors": 0,
        }
        self.cost_categories_by_name: Dict[str, str] = {}
        self.data = self.load_data_files()

    def add_issue(
        self,
        sheet: str,
        row_no: Optional[int],
        column: Optional[str],
        severity: str,
        code: str,
        message: str,
        raw_value: Optional[Any] = None,
    ) -> None:
        if severity == "warning":
            self.stats["warnings"] += 1
        elif severity == "error":
            self.stats["errors"] += 1
        self.issues.append(
            Issue(
                sheet_name=sheet,
                row_no=row_no,
                column_name=column,
                severity=severity,
                issue_code=code,
                message=message,
                raw_value=None if raw_value is None else str(raw_value),
            )
        )

    def parse_excel_date(self, raw: Any) -> Optional[date]:
        if raw is None or raw == "":
            return None
        if isinstance(raw, datetime):
            return raw.date()
        if isinstance(raw, date):
            return raw
        if isinstance(raw, (int, float)):
            try:
                return EXCEL_DATE_BASE + timedelta(days=int(raw))
            except Exception:
                return None
        if isinstance(raw, str):
            text = raw.strip()
            for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%Y.%m", "%Y/%m/%d"):
                try:
                    parsed = datetime.strptime(text, fmt)
                    return parsed.date()
                except ValueError:
                    continue
        return None

    def parse_decimal(self, raw: Any, default: Decimal = Decimal("0")) -> Optional[Decimal]:
        if raw is None or raw == "":
            return default
        if isinstance(raw, Decimal):
            return raw
        if isinstance(raw, (int, float)):
            return Decimal(str(raw))
        if isinstance(raw, str):
            text = raw.strip().replace("'", "").replace(" ", "")
            text = text.replace(",", ".")
            if text == "":
                return default
            try:
                return Decimal(text)
            except InvalidOperation:
                return None
        return None

    def normalize_text(self, raw: Any) -> str:
        if raw is None:
            return ""
        return str(raw).strip()

    def load_workbook(self) -> Any:
        return openpyxl.load_workbook(self.workbook_path, data_only=True, keep_vba=True)

    def load_json_list(self, name: str) -> List[Dict[str, Any]]:
        path = self.output_dir / name
        if not path.exists():
            return []
        with path.open("r", encoding="utf-8") as f:
            payload = json.load(f)
        if not isinstance(payload, list):
            raise RuntimeError(f"Invalid JSON format in {path}; expected a list.")
        return payload

    ARCHIVE_SUBDIR = "Archiv"
    ARCHIVE_RETENTION_DAYS = 30

    def _prune_old_archives(self) -> None:
        """Entfernt Dateien in Archiv/, deren mtime älter als ARCHIVE_RETENTION_DAYS ist."""
        archiv_dir = self.output_dir / self.ARCHIVE_SUBDIR
        if not archiv_dir.is_dir():
            return
        cutoff = time.time() - (self.ARCHIVE_RETENTION_DAYS * 86400)
        for p in archiv_dir.iterdir():
            if not p.is_file():
                continue
            try:
                if p.stat().st_mtime < cutoff:
                    p.unlink()
            except OSError:
                pass

    @staticmethod
    def _archive_filename_timestamp(now: Optional[datetime] = None) -> str:
        """YYYYMMDDHHmm, Minuten (Uhr) auf 10-Minuten-Takte gerundet (wie JsonStore)."""
        dt = now or datetime.now()
        mod = dt.hour * 60 + dt.minute + dt.second / 60.0
        rounded = int(round(mod / 10.0)) * 10
        if rounded >= 1440:
            dt = (dt + timedelta(days=1)).replace(hour=0, minute=0, second=0, microsecond=0)
            return dt.strftime("%Y%m%d") + "0000"
        h, m = divmod(rounded, 60)
        return dt.strftime("%Y%m%d") + f"{h:02d}{m:02d}"

    def _archive_existing_json_before_write(self, path: Path) -> None:
        """Kopiert eine bestehende JSON-Datei nach output_dir/Archiv (wie JsonStore in PHP)."""
        if not path.is_file():
            return
        archiv_dir = self.output_dir / self.ARCHIVE_SUBDIR
        archiv_dir.mkdir(parents=True, exist_ok=True)
        stem = path.stem
        suffix = path.suffix if path.suffix else ".json"
        stamp = self._archive_filename_timestamp()
        arch_path = archiv_dir / f"{stem}_{stamp}{suffix}"
        while arch_path.is_file():
            arch_path = archiv_dir / f"{stem}_{stamp}_{secrets.token_hex(2)}{suffix}"
        try:
            # shutil.copy (nicht copy2): neue mtime = jetzt, sonst schneidet Retention frische Archive ab.
            shutil.copy(path, arch_path)
        except OSError as exc:
            raise RuntimeError(f"Could not archive file before write: {path.name}") from exc

    def write_json_list(self, name: str, rows: List[Dict[str, Any]]) -> None:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        path = self.output_dir / name
        self._archive_existing_json_before_write(path)
        tmp = self.output_dir / f".{name}.tmp"
        with tmp.open("w", encoding="utf-8") as f:
            json.dump(rows, f, ensure_ascii=False, indent=2)
        os.replace(tmp, path)
        self._prune_old_archives()

    def load_data_files(self) -> Dict[str, List[Dict[str, Any]]]:
        return {
            "cost_categories": self.load_json_list("cost_categories.json"),
            "transactions": self.load_json_list("transactions.json"),
            "budget_plan_lines": self.load_json_list("budget_plan_lines.json"),
            "inventory_items": self.load_json_list("inventory_items.json"),
            "import_runs": self.load_json_list("import_runs.json"),
            "import_issues": self.load_json_list("import_issues.json"),
        }

    def save_data_files(self) -> None:
        self.write_json_list("cost_categories.json", self.data["cost_categories"])
        self.write_json_list("transactions.json", self.data["transactions"])
        self.write_json_list("budget_plan_lines.json", self.data["budget_plan_lines"])
        self.write_json_list("inventory_items.json", self.data["inventory_items"])
        self.write_json_list("import_runs.json", self.data["import_runs"])
        self.write_json_list("import_issues.json", self.data["import_issues"])

    def create_import_run(self) -> str:
        run_id = str(uuid4())
        self.data["import_runs"].append(
            {
                "id": run_id,
                "source_filename": os.path.basename(self.workbook_path),
                "started_at": datetime.now().isoformat(timespec="seconds"),
                "finished_at": None,
                "status": "running",
                "summary": {},
            }
        )
        return run_id

    def finalize_import_run(self, run_id: str) -> None:
        status = "success"
        if self.stats["errors"] > 0 and self.stats["warnings"] > 0:
            status = "partial"
        elif self.stats["errors"] > 0:
            status = "failed"
        elif self.stats["warnings"] > 0:
            status = "partial"
        for row in self.data["import_runs"]:
            if row["id"] == run_id:
                row["finished_at"] = datetime.now().isoformat(timespec="seconds")
                row["status"] = status
                row["summary"] = self.stats
                break

    def persist_issues(self, run_id: str) -> None:
        for issue in self.issues:
            row = asdict(issue)
            row["id"] = str(uuid4())
            row["import_run_id"] = run_id
            row["created_at"] = datetime.now().isoformat(timespec="seconds")
            self.data["import_issues"].append(row)

    def load_cost_category_map(self) -> None:
        self.cost_categories_by_name = {}
        for row in self.data["cost_categories"]:
            name = self.normalize_text(row.get("name"))
            code = self.normalize_text(row.get("code"))
            if not name or not code:
                continue
            self.cost_categories_by_name[name.lower()] = code

    @staticmethod
    def amount_to_float(raw: Optional[Decimal]) -> Optional[float]:
        if raw is None:
            return None
        return float(raw.quantize(Decimal("0.01")))

    @staticmethod
    def get_cell_by_header(ws: Any, row_no: int, header_index: Dict[str, int], header_name: str) -> Any:
        col = header_index.get(header_name)
        if not col:
            return None
        return ws.cell(row=row_no, column=col).value

    @staticmethod
    def get_header_index_with_aliases(
        ws: Any, header_row: int, max_cols: int, aliases: Dict[str, List[str]]
    ) -> Dict[str, int]:
        raw_headers: Dict[str, int] = {}
        for col in range(1, max_cols + 1):
            text = str(ws.cell(row=header_row, column=col).value or "").strip()
            if text:
                raw_headers[text] = col

        resolved: Dict[str, int] = {}
        for canonical, options in aliases.items():
            for option in options:
                if option in raw_headers:
                    resolved[canonical] = raw_headers[option]
                    break
        return resolved

    def import_cost_categories(self, ws: Any) -> None:
        by_code = {self.normalize_text(r.get("code")): r for r in self.data["cost_categories"]}
        for row_no in range(COST_CATEGORY_HEADER_ROW + 1, ws.max_row + 1):
            code = self.normalize_text(ws[f"A{row_no}"].value)
            name = self.normalize_text(ws[f"B{row_no}"].value)
            if code == "" and name == "":
                continue
            # The workbook currently has reserved KA rows (11-20) without names.
            if name == "" and code.isdigit() and int(code) >= 11:
                continue
            if code == "" or name == "":
                self.add_issue(
                    SHEET_COST_CATEGORIES,
                    row_no,
                    "A/B",
                    "warning",
                    "CATEGORY_INCOMPLETE",
                    "Kostenart row incomplete; skipping",
                    f"{code}|{name}",
                )
                continue
            item = {
                "code": code,
                "name": name,
                "is_active": True,
                "updated_at": datetime.now().isoformat(timespec="seconds"),
            }
            if code in by_code:
                by_code[code].update(item)
            else:
                item["id"] = str(uuid4())
                item["created_at"] = datetime.now().isoformat(timespec="seconds")
                self.data["cost_categories"].append(item)
                by_code[code] = item
            self.stats["cost_categories"] += 1

    def import_transactions(self, ws: Any) -> None:
        header_index = self.get_header_index_with_aliases(
            ws,
            TRANSACTION_HEADER_ROW,
            40,
            {
                "ID": ["ID"],
                "Datum": ["Datum"],
                "Empfänger": ["Empfänger", "Empf�nger"],
                "Kostenart": ["Kostenart"],
                "Konto": ["Konto"],
                "Haben/Aus": ["Haben/Aus"],
                "Soll/Ein": ["Soll/Ein"],
                "Saldo": ["Saldo"],
                "W": ["W"],
                "AG": ["AG"],
                "Bemerkung/ Beleg": ["Bemerkung/ Beleg"],
                "Status": ["Status", "Buchungsstatus", "Buchungstatus"],
            },
        )

        required_headers = ["ID", "Datum", "Empfänger", "Kostenart", "Haben/Aus", "Soll/Ein"]
        for col in required_headers:
            if col not in header_index:
                raise RuntimeError(f"Missing expected header in {SHEET_TRANSACTION}: {col}")

        by_external = {r["external_id"]: r for r in self.data["transactions"]}
        for row_no in range(TRANSACTION_HEADER_ROW + 1, ws.max_row + 1):
            external_id = self.normalize_text(self.get_cell_by_header(ws, row_no, header_index, "ID"))
            if not external_id:
                continue
            booking_date_raw = self.get_cell_by_header(ws, row_no, header_index, "Datum")
            booking_date = self.parse_excel_date(booking_date_raw)
            if not booking_date:
                self.add_issue(
                    SHEET_TRANSACTION,
                    row_no,
                    "Datum",
                    "error",
                    "INVALID_DATE",
                    "Invalid booking date",
                    booking_date_raw,
                )
                continue

            counterparty = self.normalize_text(self.get_cell_by_header(ws, row_no, header_index, "Empfänger"))
            if not counterparty:
                self.add_issue(
                    SHEET_TRANSACTION,
                    row_no,
                    "Empfänger",
                    "error",
                    "MISSING_COUNTERPARTY",
                    "Counterparty is required",
                    None,
                )
                continue

            cost_name = self.normalize_text(self.get_cell_by_header(ws, row_no, header_index, "Kostenart"))
            category_code = self.cost_categories_by_name.get(cost_name.lower()) if cost_name else None
            if cost_name and not category_code:
                self.add_issue(
                    SHEET_TRANSACTION,
                    row_no,
                    "Kostenart",
                    "warning",
                    "UNKNOWN_CATEGORY",
                    "Cost category not found; saved as raw value",
                    cost_name,
                )

            amount_out = self.parse_decimal(self.get_cell_by_header(ws, row_no, header_index, "Haben/Aus"))
            amount_in = self.parse_decimal(self.get_cell_by_header(ws, row_no, header_index, "Soll/Ein"))
            if amount_out is None or amount_in is None:
                self.add_issue(
                    SHEET_TRANSACTION,
                    row_no,
                    "Haben/Aus|Soll/Ein",
                    "error",
                    "INVALID_AMOUNT",
                    "Invalid numeric value",
                    None,
                )
                continue
            if amount_out < 0 or amount_in < 0:
                self.add_issue(
                    SHEET_TRANSACTION,
                    row_no,
                    "Haben/Aus|Soll/Ein",
                    "error",
                    "NEGATIVE_AMOUNT",
                    "Amounts must not be negative",
                    f"{amount_out}|{amount_in}",
                )
                continue
            if amount_out == 0 and amount_in == 0:
                self.add_issue(
                    SHEET_TRANSACTION,
                    row_no,
                    "Haben/Aus|Soll/Ein",
                    "error",
                    "ZERO_AMOUNT",
                    "Both amounts are zero",
                    None,
                )
                continue

            account = self.normalize_text(self.get_cell_by_header(ws, row_no, header_index, "Konto"))
            currency = self.normalize_text(self.get_cell_by_header(ws, row_no, header_index, "W")) or "CHF"
            ag_cell = self.normalize_text(self.get_cell_by_header(ws, row_no, header_index, "AG"))
            ag_lower = ag_cell.strip().lower()
            note = self.normalize_text(self.get_cell_by_header(ws, row_no, header_index, "Bemerkung/ Beleg"))
            balance_raw = self.get_cell_by_header(ws, row_no, header_index, "Saldo")
            running_balance = self.parse_decimal(balance_raw, default=Decimal("0"))
            if running_balance is None:
                running_balance = None

            item = {
                "external_id": external_id,
                "booking_date": booking_date.isoformat(),
                "counterparty": counterparty,
                "cost_category_code": category_code,
                "cost_category_name_raw": None if category_code else (cost_name or None),
                "account": account or None,
                "amount_in": self.amount_to_float(amount_in),
                "amount_out": self.amount_to_float(amount_out),
                "currency": currency[:3].upper(),
                "note": note or None,
                "running_balance_imported": self.amount_to_float(running_balance),
                "source_sheet": SHEET_TRANSACTION,
                "source_row_no": row_no,
                "updated_at": datetime.now().isoformat(timespec="seconds"),
            }
            is_new = external_id not in by_external
            booking_from_ag: Optional[str] = None
            if ag_lower == "a":
                booking_from_ag = "x"
            elif ag_lower in ("n", "v", "x"):
                booking_from_ag = ag_lower

            if "Status" in header_index:
                status_cell = self.normalize_text(
                    self.get_cell_by_header(ws, row_no, header_index, "Status")
                )
                s = status_cell.strip().lower()
                item["booking_status"] = s if s in ("n", "v", "x") else "x"
            elif booking_from_ag is not None:
                item["booking_status"] = booking_from_ag
            elif is_new:
                item["booking_status"] = "x"

            if external_id in by_external:
                row = by_external[external_id]
                row.update(item)
                row.pop("scope_code", None)
            else:
                item["id"] = str(uuid4())
                item["created_at"] = datetime.now().isoformat(timespec="seconds")
                if "booking_status" not in item:
                    item["booking_status"] = "x"
                self.data["transactions"].append(item)
                by_external[external_id] = item
            self.stats["transactions"] += 1

    def import_budget(self, ws: Any) -> None:
        header_index = self.get_header_index_with_aliases(
            ws,
            BUDGET_HEADER_ROW,
            40,
            {
                "Datum": ["Datum"],
                "Einkommen": ["Einkommen"],
                "Spenden": ["Spenden"],
                "Investitionen": ["Investitionen"],
                "Kosten": ["Kosten"],
                "Ausbildung": ["Ausbildung"],
                "Mitgliederbeiträge": ["Mitgliederbeiträge", "Mitgliederbeitr�ge"],
                "Events": ["Events"],
                "Leihabe/Kredite": ["Leihabe/Kredite"],
                "Einsätze": ["Einsätze", "Eins�tze"],
                "Depot": ["Depot"],
            },
        )
        if "Datum" not in header_index:
            raise RuntimeError(f"Missing expected header in {SHEET_BUDGET}: Datum")

        # We map every known cost category column as one budget line.
        budget_category_columns: List[Tuple[str, int]] = []
        for name, idx in header_index.items():
            if name == "Datum":
                continue
            if name.lower() in self.cost_categories_by_name:
                budget_category_columns.append((name, idx))

        if not budget_category_columns:
            self.add_issue(
                SHEET_BUDGET,
                BUDGET_HEADER_ROW,
                None,
                "warning",
                "NO_BUDGET_CATEGORY_MATCH",
                "No budget columns mapped to cost categories",
                None,
            )
            return

        by_key = {f"{r['period_date']}|{r['cost_category_code']}": r for r in self.data["budget_plan_lines"]}
        for row_no in range(BUDGET_HEADER_ROW + 1, ws.max_row + 1):
            period_raw = self.get_cell_by_header(ws, row_no, header_index, "Datum")
            period = self.parse_excel_date(period_raw)
            if not period:
                period_raw_text = self.normalize_text(period_raw).lower()
                if period_raw_text in {"summe", "total", "subtotal"}:
                    continue
                # skip completely empty rows
                row_values = [self.get_cell_by_header(ws, row_no, header_index, cat) for cat, _ in budget_category_columns]
                if all(v in (None, "") for v in row_values):
                    continue
                self.add_issue(
                    SHEET_BUDGET,
                    row_no,
                    "Datum",
                    "error",
                    "INVALID_PERIOD_DATE",
                    "Invalid budget period date",
                    period_raw,
                )
                continue
            for cat_name, col_idx in budget_category_columns:
                amount = self.parse_decimal(ws.cell(row=row_no, column=col_idx).value)
                if amount is None:
                    self.add_issue(
                        SHEET_BUDGET,
                        row_no,
                        cat_name,
                        "error",
                        "INVALID_BUDGET_AMOUNT",
                        "Invalid numeric budget value",
                        self.get_cell_by_header(ws, row_no, header_index, cat_name),
                    )
                    continue
                category_code = self.cost_categories_by_name.get(cat_name.lower())
                if not category_code:
                    self.add_issue(
                        SHEET_BUDGET,
                        row_no,
                        cat_name,
                        "warning",
                        "UNKNOWN_BUDGET_CATEGORY",
                        "Budget category missing in cost_categories",
                        cat_name,
                    )
                    continue

                key = f"{period.isoformat()}|{category_code}"
                item = {
                    "period_date": period.isoformat(),
                    "cost_category_code": category_code,
                    "amount": self.amount_to_float(amount),
                    "currency": "CHF",
                    "source_sheet": SHEET_BUDGET,
                    "source_row_no": row_no,
                    "updated_at": datetime.now().isoformat(timespec="seconds"),
                }
                if key in by_key:
                    by_key[key].update(item)
                else:
                    item["id"] = str(uuid4())
                    item["created_at"] = datetime.now().isoformat(timespec="seconds")
                    self.data["budget_plan_lines"].append(item)
                    by_key[key] = item
                self.stats["budget_lines"] += 1

    def import_inventory(self, ws: Any) -> None:
        header_index = self.get_header_index_with_aliases(
            ws,
            INVENTORY_HEADER_ROW,
            30,
            {
                "INV#": ["INV#"],
                "#Anzahl": ["#Anzahl"],
                "Artikel": ["Artikel"],
                "Beschreibung": ["Beschreibung"],
                "Kaufpreis pro Einheit": ["Kaufpreis pro Einheit"],
                "Kaufdatum": ["Kaufdatum"],
                "Neuwert": ["Neuwert"],
                "Abschreiber\n[%/Jahr]": ["Abschreiber\n[%/Jahr]"],
                "Zeitwert": ["Zeitwert"],
                "Verkauf, Liquidation": ["Verkauf, Liquidation"],
                "Verkaufswert": ["Verkaufswert"],
            },
        )
        required = ["INV#", "#Anzahl", "Artikel", "Kaufpreis pro Einheit", "Abschreiber\n[%/Jahr]"]
        for col in required:
            if col not in header_index:
                raise RuntimeError(f"Missing expected header in {SHEET_INVENTORY}: {col}")

        by_inventory = {r["inventory_no"]: r for r in self.data["inventory_items"]}
        for row_no in range(INVENTORY_HEADER_ROW + 1, ws.max_row + 1):
            inventory_no = self.normalize_text(self.get_cell_by_header(ws, row_no, header_index, "INV#"))
            if inventory_no == "":
                continue
            # "Depotzahlungen" starts a different section in this sheet, not inventory items.
            if inventory_no.lower() == "depotzahlungen":
                break

            quantity = self.parse_decimal(self.get_cell_by_header(ws, row_no, header_index, "#Anzahl"))
            item_name = self.normalize_text(self.get_cell_by_header(ws, row_no, header_index, "Artikel"))
            description = self.normalize_text(self.get_cell_by_header(ws, row_no, header_index, "Beschreibung"))
            purchase_price = self.parse_decimal(self.get_cell_by_header(ws, row_no, header_index, "Kaufpreis pro Einheit"))
            purchase_period = self.normalize_text(self.get_cell_by_header(ws, row_no, header_index, "Kaufdatum"))
            new_value = self.parse_decimal(
                self.get_cell_by_header(ws, row_no, header_index, "Neuwert"), default=Decimal("0")
            )
            depreciation = self.parse_decimal(self.get_cell_by_header(ws, row_no, header_index, "Abschreiber\n[%/Jahr]"))
            current_value = self.parse_decimal(
                self.get_cell_by_header(ws, row_no, header_index, "Zeitwert"), default=Decimal("0")
            )
            disposal_period = self.normalize_text(self.get_cell_by_header(ws, row_no, header_index, "Verkauf, Liquidation"))
            disposal_value = self.parse_decimal(
                self.get_cell_by_header(ws, row_no, header_index, "Verkaufswert"), default=Decimal("0")
            )

            if quantity is None or quantity < 0:
                self.add_issue(SHEET_INVENTORY, row_no, "#Anzahl", "error", "INVALID_QUANTITY", "Invalid quantity", None)
                continue
            if not item_name:
                self.add_issue(SHEET_INVENTORY, row_no, "Artikel", "error", "MISSING_ITEM_NAME", "Item name required", None)
                continue
            if purchase_price is None:
                self.add_issue(
                    SHEET_INVENTORY,
                    row_no,
                    "Kaufpreis pro Einheit",
                    "error",
                    "INVALID_PURCHASE_PRICE",
                    "Invalid purchase price",
                    None,
                )
                continue
            if depreciation is None or depreciation < 0 or depreciation > 100:
                self.add_issue(
                    SHEET_INVENTORY,
                    row_no,
                    "Abschreiber\n[%/Jahr]",
                    "error",
                    "INVALID_DEPRECIATION",
                    "Depreciation must be between 0 and 100",
                    None,
                )
                continue

            item = {
                "inventory_no": inventory_no,
                "quantity": self.amount_to_float(quantity),
                "item_name": item_name,
                "description": description or None,
                "purchase_price_unit": self.amount_to_float(purchase_price),
                "purchase_period": purchase_period or None,
                "new_value": self.amount_to_float(new_value),
                "depreciation_rate_pct": self.amount_to_float(depreciation),
                "current_value": self.amount_to_float(current_value),
                "disposal_period": disposal_period or None,
                "disposal_value": self.amount_to_float(disposal_value),
                "source_sheet": SHEET_INVENTORY,
                "source_row_no": row_no,
                "updated_at": datetime.now().isoformat(timespec="seconds"),
            }
            if inventory_no in by_inventory:
                by_inventory[inventory_no].update(item)
            else:
                item["id"] = str(uuid4())
                item["created_at"] = datetime.now().isoformat(timespec="seconds")
                self.data["inventory_items"].append(item)
                by_inventory[inventory_no] = item
            self.stats["inventory_items"] += 1

    def run(self) -> Dict[str, Any]:
        wb = self.load_workbook()
        required_sheets = [SHEET_COST_CATEGORIES, SHEET_TRANSACTION, SHEET_BUDGET, SHEET_INVENTORY]
        for sheet in required_sheets:
            if sheet not in wb.sheetnames:
                raise RuntimeError(f"Workbook missing required sheet: {sheet}")

        run_id = self.create_import_run()
        try:
            self.import_cost_categories(wb[SHEET_COST_CATEGORIES])
            self.load_cost_category_map()
            self.import_transactions(wb[SHEET_TRANSACTION])
            self.import_budget(wb[SHEET_BUDGET])
            self.import_inventory(wb[SHEET_INVENTORY])

            self.persist_issues(run_id)
            self.finalize_import_run(run_id)

            if not self.dry_run:
                self.save_data_files()

            return {"import_run_id": run_id, "stats": self.stats, "issues": len(self.issues)}
        except Exception:
            raise


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import Rehkitzrettung XLSM into JSON files")
    parser.add_argument("--workbook", required=True, help="Path to .xlsm workbook")
    parser.add_argument(
        "--output-dir",
        default="data",
        help="Directory where JSON files are stored (default: ./data)",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        default=False,
        help="Run full import and validation but do not write JSON files",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not os.path.exists(args.workbook):
        print(f"Workbook not found: {args.workbook}", file=sys.stderr)
        return 2

    importer = Importer(workbook_path=args.workbook, output_dir=args.output_dir, dry_run=args.dry_run)
    result = importer.run()
    print(json.dumps(result, ensure_ascii=False, indent=2))

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
